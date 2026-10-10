from docx import Document
from pathlib import Path
from .models import ExtractionResult, ExtractionMethod

def extract_docx_text(path: Path) -> ExtractionResult | None:
    document = Document(path)
    text = []

    # Paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            text.append(" | ".join(cell.text for cell in row.cells))

    if text:
        return ExtractionResult(
            text="\n".join(text),
            extraction_method=ExtractionMethod.DOCX,
        )
    else:
        return None

if __name__ == "__main__":
    file_path = Path(__file__).parent.parent.parent / "test_documents" / "this_is_docx.docx"

    extracted_text = extract_docx_text(file_path)

    print(extracted_text)